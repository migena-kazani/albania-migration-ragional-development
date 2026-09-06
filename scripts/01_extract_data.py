# -*- coding: utf-8 -*-
"""
Script 01: Data Extraction
Produces the raw regional and national panel data used throughout this replication package,
and directly generates:
  - Table 1. Net migration by region (in thousand persons), INSTAT 2024
  - Table 2. GDP by region (Gross Value Added, in billion ALL), INSTAT 2024
  - Table 3. Unemployment rate by region (%), 2016-2024, INSTAT 2024
"""
from openpyxl import load_workbook
import json

REGIONS = ['Berat', 'Dibër', 'Durrës', 'Elbasan', 'Fier', 'Gjirokastër', 'Korçë',
           'Kukës', 'Lezhë', 'Shkodër', 'Tiranë', 'Vlorë']

def extract_regional_data():
    wb = load_workbook("../data/Databaza_Scopus.xlsx", data_only=True)
    ws1 = wb['Sheet1']
    rows = list(ws1.iter_rows(min_row=1, max_row=83, values_only=True))

    # Net migration (thousand persons), 2014-2023
    net_mig = {r[0]: r[21:31] for r in rows[5:17]}
    years_net = list(range(2014, 2024))

    # GDP (Gross Value Added, million ALL), 2000-2023
    gva = {r[2]: r[3:27] for r in rows[56:68]}
    years_gva = list(range(2000, 2024))

    # Unemployment rate (%), 2016-2024
    unemp = {r[2]: [float(str(v).replace(',', '.')) if v is not None else None for v in r[3:12]]
             for r in rows[22:34]}
    years_unemp = list(range(2016, 2025))

    # Cross-sectional: population 65+, GDP per capita 2023, urban population 2011
    pop65_abs = {r[2]: r[23] for r in rows[22:34]}
    gdppc2023 = {r[5]: r[6] for r in rows[40:52]}
    urban2011 = {r[0]: r[1] for r in rows[71:83]}

    data = dict(net_mig=net_mig, years_net=years_net, gva=gva, years_gva=years_gva,
                unemp=unemp, years_unemp=years_unemp, pop65_abs=pop65_abs,
                gdppc2023=gdppc2023, urban2011=urban2011)
    with open('../output/regional_data.json', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return data


def extract_national_data():
    wb = load_workbook("../data/Databaza_Scopus.xlsx", data_only=True)
    ws3 = wb['Sheet3']
    rows = list(ws3.iter_rows(min_row=1, max_row=69, values_only=True))
    years, netmig, gdppc_lcu, unemp_nat, remit, urban_pct, urban_abs, fdi, edu = ([] for _ in range(9))
    for r in rows[4:]:
        if r[0] is None:
            continue
        years.append(int(r[0])); netmig.append(r[1]); gdppc_lcu.append(r[2])
        unemp_nat.append(r[5]); remit.append(r[6]); urban_pct.append(r[8])
        urban_abs.append(r[9]); fdi.append(r[11]); edu.append(r[13])

    data = dict(years=years, netmig=netmig, gdppc_lcu=gdppc_lcu, unemp_nat=unemp_nat,
                remit=remit, urban_pct=urban_pct, urban_abs=urban_abs, fdi=fdi, edu=edu)
    with open('../output/national_data.json', 'w') as f:
        json.dump(data, f)
    return data


def print_table1(reg_data):
    """Table 1. Net migration by region (in thousand persons), INSTAT 2024"""
    panel_years = list(range(2016, 2024))
    print("Table 1. Net migration by region (thousand persons)")
    print("Region".ljust(14), *[str(y) for y in panel_years])
    for reg in REGIONS:
        vals = [reg_data['net_mig'][reg][reg_data['years_net'].index(y)] for y in panel_years]
        print(reg.ljust(14), *["%.3f" % v for v in vals])


def print_table2(reg_data):
    """Table 2. GDP by region (Gross Value Added, in billion ALL), INSTAT 2024"""
    print("\nTable 2. GDP by region (billion ALL)")
    print("Region".ljust(14), *[str(y) for y in range(2016, 2023)])
    for reg in REGIONS:
        vals = [reg_data['gva'][reg][reg_data['years_gva'].index(y)] / 1000.0 for y in range(2016, 2023)]
        print(reg.ljust(14), *["%.1f" % v for v in vals])


def print_table3(reg_data):
    """Table 3. Unemployment rate by region (%), 2016-2024, INSTAT 2024"""
    print("\nTable 3. Unemployment rate by region (%)")
    print("Region".ljust(14), *[str(y) for y in range(2016, 2025)])
    for reg in REGIONS:
        vals = reg_data['unemp'][reg]
        print(reg.ljust(14), *["%.1f" % v for v in vals])


if __name__ == "__main__":
    reg_data = extract_regional_data()
    nat_data = extract_national_data()
    print_table1(reg_data)
    print_table2(reg_data)
    print_table3(reg_data)
    print("\nData extraction complete. Saved to ../output/regional_data.json and national_data.json")
